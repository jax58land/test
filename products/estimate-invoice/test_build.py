"""見積書・請求書Excelの数式を pycel で評価して、手計算と照合する。

usage: python test_build.py   （一時ファイルに生成して検証。通ったら python build.py で dist/ を更新）
"""
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from pycel import ExcelCompiler

from build import EXTRA, FIRST, LINES, main

EST_LAST = FIRST + LINES - 1
INV_LAST = EST_LAST + 1 + EXTRA


def build():
    path = Path(tempfile.mkdtemp()) / "t.xlsx"
    main(path)
    return path


def test_no_formula_errors():
    path = build()
    wb = load_workbook(path)
    xl = ExcelCompiler(filename=str(path))
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    v = xl.evaluate(f"'{ws.title}'!{c.coordinate}")
                    assert not (isinstance(v, str) and v.startswith("#")), f"{ws.title}!{c.coordinate} -> {v}"
                    n += 1
    assert n > 100, n


def test_estimate_totals():
    xl = ExcelCompiler(filename=str(build()))
    ev = xl.evaluate
    sub = 48000 + 2 * 18000 + 3 * 6500 + 9800 + 8000          # 121,300
    assert ev(f"見積書!G{EST_LAST + 1}") == sub
    assert ev(f"見積書!G{EST_LAST + 2}") == 12130
    assert ev(f"見積書!G{EST_LAST + 3}") == sub + 12130
    assert ev("見積書!B11") == sub + 12130
    assert ev(f"見積書!G{FIRST + 5}") == ""                    # 空行は空
    assert ev("見積書!D4") == ev("見積書!D3") + 30              # 有効期限


def test_invoice_mirrors_estimate_and_adds_extra():
    xl = ExcelCompiler(filename=str(build()))
    ev = xl.evaluate
    # 明細は見積書と同じ（転記ゼロ）
    for r in range(FIRST, FIRST + 5):
        assert ev(f"請求書!B{r}") == ev(f"見積書!B{r}")
        assert ev(f"請求書!G{r}") == ev(f"見積書!G{r}")
    assert ev(f"請求書!G{INV_LAST + 3}") == ev(f"見積書!G{EST_LAST + 3}")
    assert ev("請求書!B3") == "○○ 様"
    assert ev("請求書!F9") == "T0000000000000"                 # 登録番号（右上の発行者欄）
    # 追加工事 15,000円 を1行入れる
    xr = EST_LAST + 2
    for col in "BDF":
        ev(f"請求書!{col}{xr}")                                 # pycel: 値を入れる前に一度評価しておく
    xl.set_value(f"請求書!B{xr}", "追加工事 既設配線引替")
    xl.set_value(f"請求書!D{xr}", 1)
    xl.set_value(f"請求書!F{xr}", 15000)
    assert ev(f"請求書!G{INV_LAST + 1}") == 121300 + 15000
    assert ev(f"請求書!G{INV_LAST + 2}") == 13630
    assert ev(f"請求書!G{INV_LAST + 3}") == 136300 + 13630
    # 減額（マイナスの数量）
    xl.set_value(f"請求書!D{xr}", -1)
    assert ev(f"請求書!G{INV_LAST + 1}") == 121300 - 15000


def test_tax_rounds_down_once():
    xl = ExcelCompiler(filename=str(build()))
    xl.evaluate(f"見積書!F{FIRST}")
    xl.set_value(f"見積書!F{FIRST}", 48005)                     # 小計 121,305 → 税 12,130.5 → 12,130
    assert xl.evaluate(f"見積書!G{EST_LAST + 1}") == 121305
    assert xl.evaluate(f"見積書!G{EST_LAST + 2}") == 12130
    assert xl.evaluate(f"請求書!G{INV_LAST + 2}") == 12130      # 請求書にも反映


def test_invoice_required_fields_present():
    """インボイスの記載事項に対応する欄がそろっていること。"""
    wb = load_workbook(build())
    inv = wb["請求書"]
    labels = {c.value for row in inv.iter_rows() for c in row if isinstance(c.value, str)}
    for need in ["登録番号", "取引年月日", "10%対象 小計", "消費税（10%）", "宛先", "品名・工事内容"]:
        assert need in labels, need
    guide = " ".join(str(c.value) for c in wb["使い方"]["A"] if c.value)
    assert "税理士" in guide and "国税庁" in guide


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn()
            print("ok", name)
    print("ALL PASSED")
