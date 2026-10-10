"""電気工事 見積書・請求書 Excel（インボイス対応・単独版）を生成する。

usage: python build.py [出力パス]   （test_build.py は一時ファイルに作って検証する）
出力先: リポジトリ直下の dist/denki-estimate-invoice.xlsx
- 自社情報（屋号・登録番号・振込先）は「自社情報」シートに1回入れるだけ
- 請求書の明細は見積書を参照する（転記ゼロ）。追加・変更工事は請求書側の追加欄に入れる
- 消費税は請求書1枚につき1回だけ端数処理（切り捨て）
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

VERSION = "1.0"
FONT = "Meiryo UI"
DIST = Path(__file__).resolve().parents[2] / "dist"
OUT_NAME = "denki-estimate-invoice.xlsx"

F_BASE = Font(name=FONT, size=10)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=14, bold=True)
F_DOC = Font(name=FONT, size=18, bold=True)
F_BIG = Font(name=FONT, size=16, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_NOTE = Font(name=FONT, size=9, color="666666")
FILL_INPUT = PatternFill("solid", fgColor="FFF9C4")
FILL_HEAD = PatternFill("solid", fgColor="37474F")
FILL_OUT = PatternFill("solid", fgColor="E3F2FD")
THIN = Side(style="thin", color="B0BEC5")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(vertical="top", wrap_text=True)
YEN = "#,##0"

DISCLAIMER = ("【免責】本ファイルは見積書・請求書の作成を補助するテンプレートです。インボイス制度・消費税の取扱いの最終判断は、"
              "国税庁の公表資料や税理士・税務署にご確認ください。")

LINES = 20      # 見積明細の行数
EXTRA = 5       # 請求書の追加・変更工事の行数
FIRST = 14      # 明細の開始行（見積書・請求書とも）


def cell(ws, ref, value=None, font=F_BASE, fill=None, fmt=None, align=None, border=False):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = align
    if border:
        c.border = BORDER
    return c


def inp(ws, ref, value=None, fmt=None, border=False):
    return cell(ws, ref, value, F_INPUT, FILL_INPUT, fmt, border=border)


def out(ws, ref, formula, fmt=None, font=F_BASE, border=False, align=None):
    return cell(ws, ref, formula, font, FILL_OUT, fmt, align, border)


def header_row(ws, row, headers):
    for i, h in enumerate(headers):
        cell(ws, f"{chr(ord('A') + i)}{row}", h, F_HEAD, FILL_HEAD, align=CENTER, border=True)
    ws.row_dimensions[row].height = 26


# ---------------------------------------------------------------- 自社情報
COMPANY = [  # (行, 項目, 例)
    (4, "屋号・会社名", "○○電設"),
    (5, "担当者", "○○"),
    (6, "住所", "○○県○○市○○ 1-2-3"),
    (7, "電話", "000-0000-0000"),
    (8, "登録番号（インボイス）", "T0000000000000"),
    (10, "振込先 銀行・支店", "○○銀行 ○○支店"),
    (11, "振込先 種別・口座番号", "普通 0000000"),
    (12, "振込先 口座名義", "マルマル デンセツ"),
    (14, "見積の有効期限（日数）", 30),
    (15, "支払期限（請求日から日数）", 30),
    (16, "消費税率", 0.1),
]


def build_company(wb):
    ws = wb.create_sheet("自社情報")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 40
    cell(ws, "A1", "自社情報（ここに1回入れるだけで、見積書・請求書の両方に反映されます）", F_TITLE)
    cell(ws, "A2", "黄色セル（青字）＝入力欄", F_NOTE)
    for r, label, ex in COMPANY:
        cell(ws, f"A{r}", label, F_BOLD, border=True)
        inp(ws, f"B{r}", ex, "0%" if label == "消費税率" else None, border=True)
    ws["B8"].comment = Comment("適格請求書発行事業者の登録番号（T＋13桁）。未登録なら空欄に", "tool")
    ws["B16"].comment = Comment("電気工事は通常10%（軽減税率の対象外）", "tool")
    cell(ws, "A18", "※ 個人の方は、屋号がなければ氏名でもかまいません（実名を出したくない場合の扱いは税務署等にご確認ください）。", F_NOTE)
    return ws


def C(row):
    return f"自社情報!$B${row}"


def issuer_block(ws, top=5):
    """右上の発行者欄（自社情報を参照）。"""
    labels = [("発行者", 4), ("担当", 5), ("住所", 6), ("電話", 7), ("登録番号", 8)]
    for i, (lab, src) in enumerate(labels):
        cell(ws, f"E{top + i}", lab, F_BOLD)
        out(ws, f"F{top + i}", f'=IF({C(src)}="","",{C(src)})')
        ws.merge_cells(f"F{top + i}:G{top + i}")


def items_table(ws, first, n, src=None, start_no=1):
    """明細表。src を渡すと、その見積書の行を参照する（請求書用）。"""
    for i in range(n):
        r = first + i
        cell(ws, f"A{r}", start_no + i, border=True, align=CENTER)
        for col in "BCDEF":
            if src:
                ref = f"{src}!{col}{r}"
                out(ws, f"{col}{r}", f'=IF({ref}="","",{ref})', YEN if col == "F" else None, border=True)
            else:
                inp(ws, f"{col}{r}", None, YEN if col == "F" else None, border=True)
        out(ws, f"G{r}", f'=IF(OR(D{r}="",F{r}=""),"",ROUND(D{r}*F{r},0))', YEN, border=True)


def totals(ws, first, last, row):
    """小計・消費税（1回だけ切り捨て）・合計。戻り値: 合計セル。"""
    cell(ws, f"E{row}", "10%対象 小計", F_BOLD, border=True)
    ws.merge_cells(f"E{row}:F{row}")
    out(ws, f"G{row}", f"=SUM(G{first}:G{last})", YEN, F_BOLD, True)
    cell(ws, f"E{row + 1}", "消費税（10%）", F_BOLD, border=True)
    ws.merge_cells(f"E{row + 1}:F{row + 1}")
    out(ws, f"G{row + 1}", f"=ROUNDDOWN(G{row}*{C(16)},0)", YEN, F_BOLD, True)
    cell(ws, f"E{row + 2}", "合計（税込）", F_BOLD, border=True)
    ws.merge_cells(f"E{row + 2}:F{row + 2}")
    out(ws, f"G{row + 2}", f"=G{row}+G{row + 1}", YEN, F_BOLD, True)
    return f"G{row + 2}"


def page_setup(ws, last_row):
    for col, w in zip("ABCDEFG", [5, 30, 16, 8, 7, 12, 14]):
        ws.column_dimensions[col].width = w
    ws.print_area = f"A1:G{last_row}"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# ---------------------------------------------------------------- 見積書
EXAMPLES = [
    ("分電盤交換", "主幹50A 16回路", 1, "台", 48000),
    ("エアコン専用回路増設", "VVF2.0-3C 15m程度", 2, "回路", 18000),
    ("コンセント増設", "2口 接地付", 3, "箇所", 6500),
    ("材料・消耗品", "", 1, "式", 9800),
    ("諸経費", "", 1, "式", 8000),
]


def build_estimate(wb):
    ws = wb.create_sheet("見積書")
    ws.merge_cells("A1:G1")
    cell(ws, "A1", "御 見 積 書", F_DOC, align=CENTER)
    for r, lab, v, fmt in [(3, "宛先", "○○ 様", None), (4, "件名", "○○邸 電気設備工事", None),
                           (5, "工事場所", "○○市○○", None), (6, "見積番号", "M-0001", None)]:
        cell(ws, f"A{r}", lab, F_BOLD)
        inp(ws, f"B{r}", v, fmt)
    cell(ws, "C3", "見積日", F_BOLD)
    inp(ws, "D3", "=TODAY()", "yyyy/mm/dd")
    cell(ws, "C4", "有効期限", F_BOLD)
    out(ws, "D4", f'=IF(D3="","",D3+{C(14)})', "yyyy/mm/dd")
    issuer_block(ws)
    last = FIRST + LINES - 1
    cell(ws, "A11", "御見積金額（税込）", F_BOLD)
    header_row(ws, FIRST - 1, ["No", "品名・工事内容", "規格・内容", "数量", "単位", "単価", "金額"])
    items_table(ws, FIRST, LINES)
    for i, ex in enumerate(EXAMPLES):
        for col, v in zip("BCDEF", ex):
            ws[f"{col}{FIRST + i}"].value = v
    total = totals(ws, FIRST, last, last + 1)
    out(ws, "B11", f"={total}", '"¥"#,##0"-"', F_BIG)
    cell(ws, f"A{last + 5}", "備考", F_BOLD)
    ws.merge_cells(f"B{last + 5}:G{last + 7}")
    inp(ws, f"B{last + 5}", "・既設配線の不良等により追加工事が必要な場合は、事前にご相談のうえ別途お見積りいたします。")
    ws[f"B{last + 5}"].alignment = WRAP
    page_setup(ws, last + 7)
    return ws


# ---------------------------------------------------------------- 請求書
def build_invoice(wb):
    ws = wb.create_sheet("請求書")
    ws.merge_cells("A1:G1")
    cell(ws, "A1", "請 求 書", F_DOC, align=CENTER)
    cell(ws, "A3", "宛先", F_BOLD)
    out(ws, "B3", '=見積書!B3')
    cell(ws, "A4", "件名", F_BOLD)
    out(ws, "B4", '=見積書!B4')
    cell(ws, "A5", "工事場所", F_BOLD)
    out(ws, "B5", '=見積書!B5')
    cell(ws, "A6", "請求番号", F_BOLD)
    inp(ws, "B6", "S-0001")
    cell(ws, "C3", "請求日", F_BOLD)
    inp(ws, "D3", "=TODAY()", "yyyy/mm/dd")
    cell(ws, "C4", "取引年月日", F_BOLD)
    inp(ws, "D4", None, "yyyy/mm/dd")
    ws["C4"].comment = Comment("工事の完了日（引渡し日）を入れてください。インボイスの記載事項です", "tool")
    cell(ws, "C5", "お支払期限", F_BOLD)
    out(ws, "D5", f'=IF(D3="","",D3+{C(15)})', "yyyy/mm/dd")
    issuer_block(ws)
    est_last = FIRST + LINES - 1
    cell(ws, "A11", "ご請求金額（税込）", F_BOLD)
    header_row(ws, FIRST - 1, ["No", "品名・工事内容", "規格・内容", "数量", "単位", "単価", "金額"])
    items_table(ws, FIRST, LINES, src="見積書")
    # 追加・変更工事（請求書でだけ入力）
    xr = est_last + 1
    cell(ws, f"A{xr}", "追加・変更工事（見積後に決まったもの。減額はマイナスの数量で）", F_BOLD)
    ws.merge_cells(f"A{xr}:G{xr}")
    items_table(ws, xr + 1, EXTRA, start_no=LINES + 1)
    last = xr + EXTRA
    total = totals(ws, FIRST, last, last + 1)
    out(ws, "B11", f"={total}", '"¥"#,##0"-"', F_BIG)
    pr = last + 5
    cell(ws, f"A{pr}", "お振込先", F_BOLD)
    out(ws, f"B{pr}", f'={C(10)}&"　"&{C(11)}')
    ws.merge_cells(f"B{pr}:G{pr}")
    out(ws, f"B{pr + 1}", f'="口座名義："&{C(12)}')
    ws.merge_cells(f"B{pr + 1}:G{pr + 1}")
    cell(ws, f"A{pr + 2}", "※振込手数料はご負担くださいますようお願いいたします。", F_NOTE)
    cell(ws, f"A{pr + 3}", "※金額はすべて税率10%の対象です（軽減税率の対象品目はありません）。", F_NOTE)
    page_setup(ws, pr + 3)
    return ws


# ---------------------------------------------------------------- 使い方
def build_guide(ws):
    ws.title = "使い方"
    ws.column_dimensions["A"].width = 110
    cell(ws, "A1", f"電気工事 見積書・請求書（インボイス対応）  v{VERSION}", F_TITLE)
    lines = [
        "■ 使い方（3ステップ）",
        "1. 「自社情報」に屋号・登録番号・振込先を1回だけ入れる",
        "2. 「見積書」に宛先・件名・明細を入れて印刷（またはPDF保存）",
        "3. 工事が終わったら「請求書」の請求日・取引年月日を入れる。明細は見積書から自動で入ります。",
        "   見積後に決まった追加・変更工事だけ、請求書の「追加・変更工事」欄に入れてください（減額はマイナスの数量で）。",
        "",
        "■ 色の意味",
        "・黄色セル（青字）＝入力欄 ／ 水色セル＝自動（見積書・自社情報を参照しています。上書きしないでください）",
        "",
        "■ インボイス（適格請求書）の記載事項との対応（国税庁「適格請求書等保存方式の概要」の記載事項）",
        "・発行者の氏名または名称・登録番号 …… 自社情報 → 請求書の右上",
        "・取引年月日 …… 請求書の「取引年月日」（工事の完了日）",
        "・取引内容 …… 明細の品名・工事内容",
        "・税率ごとに区分した合計額と適用税率 …… 「10%対象 小計」",
        "・税率ごとに区分した消費税額 …… 「消費税（10%）」（請求書1枚につき1回だけ切り捨て）",
        "・交付を受ける事業者の名称 …… 宛先",
        "",
        "■ 計算",
        "・金額 = 数量 × 単価（円未満四捨五入）／ 消費税 = 小計 × 税率（円未満切り捨て、請求書1枚につき1回）",
        "・見積の有効期限・支払期限は、自社情報の日数から自動で計算します。",
        "",
        DISCLAIMER,
    ]
    for i, t in enumerate(lines):
        cell(ws, f"A{3 + i}", t, F_BOLD if t.startswith("■") else F_BASE, align=WRAP)


def main(out_path=None):
    wb = Workbook()
    build_guide(wb.active)
    build_company(wb)
    build_estimate(wb)
    build_invoice(wb)
    wb.calculation.fullCalcOnLoad = True
    out_path = Path(out_path or DIST / OUT_NAME)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(out_path)
    return out_path


if __name__ == "__main__":
    import sys
    main(sys.argv[1] if len(sys.argv) > 1 else None)
