"""生成したExcelの数式を pycel で評価し、手計算値と照合する。

usage: python test_build.py   （事前に build.py と build.py --lite を実行）
"""
from pathlib import Path

from openpyxl import load_workbook
from pycel import ExcelCompiler

DIST = Path(__file__).resolve().parents[2] / "dist"


def approx(a, b, tol=1e-6):
    return abs(float(a) - float(b)) < tol


def check_no_errors(path):
    wb = load_workbook(path)
    xl = ExcelCompiler(filename=str(path))
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    v = xl.evaluate(f"'{ws.title}'!{c.coordinate}")
                    assert not (isinstance(v, str) and v.startswith("#")), f"{ws.title}!{c.coordinate} -> {v}"
                    n += 1
    return xl, n


def test_full():
    xl, n = check_no_errors(DIST / "denki-calc-kit.xlsx")
    ev = lambda ref: xl.evaluate(ref)
    # 電圧降下: 単相2線 12A 35m 3.5mm² → 35.6*35*12/3500
    assert approx(ev("電圧降下!H6"), 35.6)
    assert approx(ev("電圧降下!I6"), 35.6 * 35 * 12 / 3500)
    assert ev("電圧降下!K6") == "NG"
    assert approx(ev("電圧降下!L6"), 35.6 * 35 * 12 / 2000)
    assert ev("電圧降下!M6") == 8
    assert ev("電圧降下!I7") == ""
    # 許容電流: 2.0mm 35A × 0.70 = 24.5 → 7捨8入で24A
    assert ev("許容電流!I6") == 24
    assert ev("許容電流!J6") == "OK"
    # 材料拾い: (18+6+22+5)*1.1=56.1→57 ×185
    assert ev("材料拾い!M6") == 57
    assert ev("材料拾い!O6") == 57 * 185
    assert ev("材料拾い!M10") == 9  # 8m ×1.05 = 8.4 → 9
    total = 57 * 185 + 74 * 95 + 5 * 180 + 10 * 420 + 9 * 60
    assert ev("材料拾い!O36") == total
    # 見積書
    sub = total + 62500 + 18000 + 8000
    assert ev("見積書!G32") == sub
    assert ev("見積書!G33") == sub // 10
    assert ev("見積書!G34") == sub + sub // 10
    print(f"full: {n} formulas OK")


def test_rounding_edges():
    """7捨8入と浮動小数の境界（手元で関数を再現して確認）"""
    def rule(x):  # シートの式と同じロジック
        import math
        return math.floor(math.floor(round(x, 6) * 10) / 10 + 0.2 + 1e-9)
    assert rule(27 * 0.70) == 19   # 18.9 → 19（電技解釈の代表例）
    assert rule(35 * 0.70) == 24   # 24.5 → 24
    assert rule(48 * 0.70) == 33   # 33.6 → 33
    assert rule(62 * 0.70) == 43   # 43.4 → 43
    assert rule(37 * 0.70) == 26   # 25.9 → 26
    assert rule(49 * 0.56) == 27   # 27.44 → 27
    print("rounding OK")


def test_lite():
    _, n = check_no_errors(DIST / "denki-calc-kit-lite.xlsx")
    print(f"lite: {n} formulas OK")


if __name__ == "__main__":
    test_full()
    test_lite()
    test_rounding_edges()
