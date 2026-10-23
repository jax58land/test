"""dist/ の販売ファイルが、いまのソースから作り直したものと同じかを確かめる（更新忘れの検出）。

usage: python tools/test_dist_fresh.py
- Excel: 全シート・全セルの値（数式）を比較
- PDF: 全ページの文字を比較
差があれば「build.py → test_build.py → dist 更新」をやり直す合図。
"""
import importlib.util
import sys
import tempfile
from pathlib import Path

import pymupdf
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DIST = ROOT / "dist"


def load(name, folder):
    sys.path.insert(0, str(folder))
    try:
        spec = importlib.util.spec_from_file_location(name, folder / "build.py")
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        return mod
    finally:
        sys.path.remove(str(folder))


def same_xlsx(a, b):
    wa, wb = load_workbook(a), load_workbook(b)
    assert wa.sheetnames == wb.sheetnames, (a.name, wa.sheetnames, wb.sheetnames)
    for ws in wa.worksheets:
        other = wb[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if c.coordinate in ("F3", "D3") and isinstance(c.value, str) and "TODAY" in c.value:
                    continue
                assert c.value == other[c.coordinate].value, f"{b.name} {ws.title}!{c.coordinate}: {other[c.coordinate].value!r} → 作り直すと {c.value!r}"


def same_pdf(a, b):
    da, db = pymupdf.open(a), pymupdf.open(b)
    assert len(da) == len(db), (b.name, len(da), len(db))
    for i in range(len(da)):
        assert da[i].get_text() == db[i].get_text(), f"{b.name} {i + 1}ページ目の文字が違う"


def test_dist_is_fresh():
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        kit = load("kit_build", ROOT / "products" / "denki-calc-kit")
        kit.DIST = tmp
        for argv, name in ((["build.py"], "denki-calc-kit.xlsx"), (["build.py", "--lite"], "denki-calc-kit-lite.xlsx")):
            sys.argv = argv
            kit.main()
            same_xlsx(tmp / name, DIST / name)
        inv = load("inv_build", ROOT / "products" / "estimate-invoice")
        same_xlsx(inv.main(tmp / "inv.xlsx"), DIST / "denki-estimate-invoice.xlsx")
        drill = load("drill_build", ROOT / "products" / "fukusenzu-drill")
        same_pdf(drill.main(tmp / "drill.pdf"), DIST / "fukusenzu-drill.pdf")
        same_pdf(drill.main(tmp / "drill-lite.pdf", lite=True), DIST / "fukusenzu-drill-lite.pdf")


if __name__ == "__main__":
    test_dist_is_fresh()
    print("ALL PASSED")
