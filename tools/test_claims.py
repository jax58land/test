"""記事・出品文に書いた数字が、実際の商品と一致しているかを照合する。

usage: python tools/test_claims.py
商品を直したのに説明文だけ古いまま、という食い違いを防ぐ。数字は販売ファイル（dist/）と問題データから数える。
"""
import sys
from pathlib import Path

import pymupdf
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DIST = ROOT / "dist"
sys.path.insert(0, str(ROOT / "products" / "fukusenzu-drill"))
from circuits import PROBLEMS  # noqa: E402


def formula_rows(ws, col):
    return sum(1 for r in range(1, ws.max_row + 1)
               if isinstance(ws[f"{col}{r}"].value, str) and ws[f"{col}{r}"].value.startswith("="))


def facts():
    kit = load_workbook(DIST / "denki-calc-kit.xlsx")
    take = kit["材料拾い"]
    heads = [c.value for c in take[5]]
    zones = heads.index("合計数量") - heads.index("ロス率") - 1      # ロス率と合計数量のあいだの列＝区画
    inv = load_workbook(DIST / "denki-estimate-invoice.xlsx")
    sys.path.insert(0, str(ROOT / "products" / "estimate-invoice"))
    from build import EXTRA, LINES
    return {
        "電圧降下の回路数": formula_rows(kit["電圧降下"], "I"),
        "許容電流の回路数": formula_rows(kit["許容電流"], "I"),
        "材料拾いの区画数": zones,
        "材料拾いの品目数": formula_rows(take, "M"),
        "ドリルの問題数": len(PROBLEMS),
        "ドリル製品版のページ数": len(pymupdf.open(DIST / "fukusenzu-drill.pdf")),
        "ドリル無料版のページ数": len(pymupdf.open(DIST / "fukusenzu-drill-lite.pdf")),
        "見積の明細行数": LINES,
        "請求の追加行数": EXTRA,
        "見積の明細行数（実ファイル）": formula_rows(inv["見積書"], "G") - 3,   # 小計・税・合計の3つを除く
    }


# (ファイル, 本文に含まれる文字列, 事実の名前, 文字列が表す数)
CLAIMS = [
    ("content/note/01-free-voltage-drop.md", "20回路まで", "電圧降下の回路数", 20),
    ("content/note/03-free-ampacity.md", "15回路まで", "許容電流の回路数", 15),
    ("content/note/05-free-estimate.md", "8区画×30品目", "材料拾いの区画数", 8),
    ("content/note/05-free-estimate.md", "8区画×30品目", "材料拾いの品目数", 30),
    ("content/x/posts.md", "20回路まで", "電圧降下の回路数", 20),
    ("content/booth/fukusenzu-drill.md", "複線図ドリル10問", "ドリルの問題数", 10),
    ("content/booth/fukusenzu-drill.md", "25ページ", "ドリル製品版のページ数", 25),
    ("content/booth/fukusenzu-drill.md", "7ページ", "ドリル無料版のページ数", 7),
    ("content/note/04-free-fukusenzu.md", "全10問", "ドリルの問題数", 10),
    ("content/booth/denki-estimate-invoice.md", "明細20行", "見積の明細行数", 20),
    ("content/booth/denki-estimate-invoice.md", "明細20行", "見積の明細行数（実ファイル）", 20),
    ("content/booth/denki-estimate-invoice.md", "5行の追加欄", "請求の追加行数", 5),
]


def test_claims_match_products():
    f = facts()
    for path, text, key, n in CLAIMS:
        body = (ROOT / path).read_text(encoding="utf-8")
        assert text in body, f"{path}: 「{text}」が見つからない（文言を変えたらこの表も直す）"
        assert f[key] == n, f"{path}: 「{text}」と書いているが、実際の{key}は {f[key]}"


if __name__ == "__main__":
    for k, v in facts().items():
        print(f"  {k}: {v}")
    test_claims_match_products()
    print("ALL PASSED")
